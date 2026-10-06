#!/usr/bin/env python3
"""
Figma PPT Insert - Marker Image Insertion

Insert Figma-exported screen images into a PowerPoint template. Targets are
found per slide, in this order:
  1. marker shapes whose text is ``[화면 교체 위치: {화면ID}]`` (or pictures this
     tool inserted earlier, which keep their screen ID and original box);
  2. otherwise, a table cell "화면ID" gives the ID and the slide's largest
     picture (the existing screenshot) is replaced in its box and z-order.
Each image is matched by file name (= Figma frame name = screen ID) and fitted
into the target box; the marker / old picture is removed.

All geometry is computed in inches; image pixels are used only for the
unitless aspect ratio.

Project layout (one folder per project under <repo>/projects/):
    projects/{project}/01_template/                one source .pptx
    projects/{project}/02_images/YYYY-MM-DD/       exports per date; newest folder wins per ID
    projects/{project}/03_output/{name}_최종.pptx  latest result
    projects/{project}/03_output/_archive/         previous results (newest 10 kept)
    projects/{project}/insert_log.md               cumulative run log

Usage:
    python3 .claude/skills/figma-ppt-insert/scripts/insert.py [options]

Examples:
    python3 .claude/skills/figma-ppt-insert/scripts/insert.py --new-project IMS
    python3 .claude/skills/figma-ppt-insert/scripts/insert.py --project IMS
    python3 .claude/skills/figma-ppt-insert/scripts/insert.py --project IMS --mode fill
    python3 .claude/skills/figma-ppt-insert/scripts/insert.py --project IMS --dry-run
    python3 .claude/skills/figma-ppt-insert/scripts/insert.py --list

Dependencies:
    python-pptx, Pillow
"""
from __future__ import annotations

import argparse
import copy
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterator, Optional

# scripts/insert.py: [0]=scripts [1]=figma-ppt-insert [2]=skills [3]=.claude [4]=repo
REPO_ROOT = Path(__file__).resolve().parents[4]

try:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Emu, Inches, Pt
except ImportError as exc:  # pragma: no cover - environment guard
    print(
        f"Missing dependency: {exc.name}. Install with: pip install python-pptx Pillow",
        file=sys.stderr,
    )
    raise SystemExit(1)

PROJECTS_DIR = REPO_ROOT / "projects"
TEMPLATE_DIR = "01_template"
IMAGES_DIR = "02_images"
OUTPUT_DIR = "03_output"
ARCHIVE_DIR = "_archive"
LOG_FILE = "insert_log.md"
OUTPUT_SUFFIX = "_최종"
ARCHIVE_KEEP = 10
IMAGE_EXTS = (".png", ".jpg", ".jpeg")  # earlier = preferred on a same-folder tie
IMAGE_FORMATS = {"PNG", "JPEG"}  # actual content, checked with Pillow
IGNORED_FILES = {".gitkeep", "thumbs.db", "desktop.ini", ".ds_store"}
MODE_FIT = "fit"
MODE_FILL = "fill"
# Font for text this tool adds to slides (install-local font lock).
FONT_FAMILY = "Pretendard"

TABLE_ID_LABEL = "화면ID"  # compared with whitespace removed, upper-cased
MIN_SCREEN_AREA_SQIN = 1.0  # smaller pictures (logos, icons) are never screen targets
SIMILAR_RATIO_TOLERANCE = 0.15  # e.g. 16:9 vs 16:10 (11%) counts as similar
PLACE_TOP = "상단 맞춤"
PLACE_CENTER = "가운데"
PLACE_FILL = "채우기"
MIN_PPI = 150  # below this, screen text looks blurry when shown on the slide
PAGE_LABEL = "페이지"
UNKNOWN_FIELD_LABELS = {"화면명", "화면타입", "유형", "화면경로"}  # left as placeholders on new pages
PLACEHOLDER = "(확인 필요)"
TYPE_LABEL = "유형"
CLONE_TYPE_VALUE = "변경"  # new pages copy the frame of the nearest slide with this 유형
DESCRIPTION_PLACEHOLDER = (
    "[확인 필요] 화면 설명",
    "피그마 export로 추가된 신규 화면입니다.",
    "화면명·화면타입·유형·화면 경로와 기능 설명을 입력하세요.",
)
FOOTER_ZONE = 0.9  # shapes below this share of the slide height are footer (page number)
FRAME_MIN_SHARE = 0.4  # empty rectangle covering this share of the slide = content frame
PIC_NAME_PREFIX = "FIGMA:"
REGION_TAG = "figma-ppt-insert region_in="

_MARKER_RE = re.compile(r"\[\s*화면\s*교체\s*위치\s*[:：]\s*(?P<id>[^\]\s]+)\s*\]")
_DATE_DIR_RE = re.compile(r"\d{4}-\d{2}-\d{2}")
_SCALE_SUFFIX_RE = re.compile(r"@\d+(?:\.\d+)?x$", re.IGNORECASE)
_BAD_PROJECT_CHARS_RE = re.compile(r'[\\/:*?"<>|]')
_NUMBER_SPLIT_RE = re.compile(r"(\d+)")
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REGION_RE = re.compile(re.escape(REGION_TAG) + r"([-\d.]+),([-\d.]+),([-\d.]+),([-\d.]+)")


@dataclass
class Box:
    """Rectangle in inches."""

    left: float
    top: float
    width: float
    height: float


@dataclass
class Target:
    slide_no: int
    screen_id: str
    kind: str  # "marker" | "picture" (inserted earlier) | "table" (화면ID table + screenshot)
    box: Box
    element: object  # lxml element of the shape to replace


@dataclass
class ImageEntry:
    screen_id: str
    path: Path
    folder: str


@dataclass
class RunReport:
    inserted: list[tuple[Target, ImageEntry, str]] = field(default_factory=list)
    missing_image: list[Target] = field(default_factory=list)
    kept_existing: list[Target] = field(default_factory=list)
    unused_images: list[ImageEntry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _key(screen_id: str) -> str:
    return screen_id.strip().upper()


def _emu_to_in(value: int) -> float:
    return Emu(value).inches


def _shape_box(shape) -> Box:
    return Box(
        _emu_to_in(shape.left),
        _emu_to_in(shape.top),
        _emu_to_in(shape.width),
        _emu_to_in(shape.height),
    )


def _marker_id(shape) -> Optional[str]:
    if not getattr(shape, "has_text_frame", False):
        return None
    match = _MARKER_RE.fullmatch(shape.text_frame.text.strip())
    return match.group("id") if match else None


def _iter_group_markers(group) -> Iterator[str]:
    for shape in group.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_group_markers(shape)
        else:
            screen_id = _marker_id(shape)
            if screen_id:
                yield screen_id


def _picture_region(shape) -> Box:
    descr = shape._element.nvPicPr.cNvPr.get("descr", "")
    match = _REGION_RE.search(descr)
    if match:
        return Box(*(float(v) for v in match.groups()))
    return _shape_box(shape)


def _table_screen_id(slide) -> Optional[str]:
    """Return the value next to a "화면ID" header cell in any top-level table."""
    for shape in slide.shapes:
        if not getattr(shape, "has_table", False) or not shape.has_table:
            continue
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            for idx, text in enumerate(cells):
                if re.sub(r"\s+", "", text).upper() != TABLE_ID_LABEL:
                    continue
                # Merged cells repeat or blank out; take the first non-empty value after it.
                for value in cells[idx + 1:]:
                    if value and re.sub(r"\s+", "", value).upper() != TABLE_ID_LABEL:
                        return value.split()[0]
    return None


def _largest_picture(slide):
    pictures = [
        s for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
        and _emu_to_in(s.width) * _emu_to_in(s.height) >= MIN_SCREEN_AREA_SQIN
    ]
    return max(pictures, key=lambda s: s.width * s.height, default=None)


def find_targets(prs, report: RunReport) -> list[Target]:
    """Collect targets slide by slide.

    A slide with markers or previously inserted pictures uses those. Otherwise,
    when a table carries a "화면ID" cell, the slide's largest picture (the
    existing screenshot) becomes the target and keeps its box and z-order.
    """
    targets: list[Target] = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        before = len(targets)
        has_group_marker = _scan_slide(slide, slide_no, targets, report)
        if has_group_marker or len(targets) > before:
            continue
        screen_id = _table_screen_id(slide)
        if not screen_id:
            continue
        picture = _largest_picture(slide)
        if picture is None:
            report.warnings.append(
                f"슬라이드 {slide_no}: 화면ID `{screen_id}`는 있지만 교체할 화면 이미지가 없습니다 — "
                "마커 도형을 넣어 주세요."
            )
            continue
        targets.append(Target(slide_no, screen_id, "table", _shape_box(picture), picture._element))
    return targets


def _scan_slide(slide, slide_no: int, targets: list[Target], report: RunReport) -> bool:
    """Append marker / inserted-picture targets; return True if a group marker was seen."""
    group_marker = False
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for screen_id in _iter_group_markers(shape):
                group_marker = True
                report.warnings.append(
                    f"슬라이드 {slide_no}: 그룹 안의 마커 `{screen_id}`는 1차 미지원 — "
                    "그룹을 해제한 뒤 다시 실행하세요."
                )
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name.startswith(PIC_NAME_PREFIX):
            screen_id = shape.name[len(PIC_NAME_PREFIX):]
            targets.append(
                Target(slide_no, screen_id, "picture", _picture_region(shape), shape._element)
            )
            continue
        screen_id = _marker_id(shape)
        if screen_id:
            if shape.rotation:
                report.warnings.append(
                    f"슬라이드 {slide_no}: 마커 `{screen_id}`가 회전돼 있습니다 — "
                    "회전 없이 영역만 사용합니다."
                )
            targets.append(Target(slide_no, screen_id, "marker", _shape_box(shape), shape._element))
    return group_marker


def _is_ignorable(path: Path) -> bool:
    name = path.name.lower()
    return name in IGNORED_FILES or name.startswith(("~$", "._"))


def _unsupported_reason(path: Path) -> Optional[str]:
    """Return why the file cannot be inserted, or None when it is a usable PNG/JPEG."""
    ext = path.suffix.lower()
    if ext not in IMAGE_EXTS:
        return f"지원하지 않는 형식({ext or '확장자 없음'})"
    try:
        with Image.open(path) as img:
            actual = img.format
    except (OSError, ValueError):
        return "이미지를 열 수 없음(손상된 파일)"
    if actual not in IMAGE_FORMATS:
        return f"확장자는 {ext}이지만 실제 형식은 {actual}"
    return None


def collect_images(images_root: Path, report: RunReport) -> dict[str, ImageEntry]:
    """Return the newest image per screen ID across YYYY-MM-DD folders."""
    dated: list[tuple[date, Path]] = []
    if images_root.is_dir():
        for folder in images_root.iterdir():
            if not folder.is_dir():
                continue
            try:
                if not _DATE_DIR_RE.fullmatch(folder.name):
                    raise ValueError
                dated.append((date.fromisoformat(folder.name), folder))
            except ValueError:
                report.warnings.append(
                    f"`{IMAGES_DIR}/{folder.name}` 폴더는 날짜 형식(YYYY-MM-DD)이 아니라 건너뜁니다."
                )
        for loose in images_root.iterdir():
            if loose.is_file() and not _is_ignorable(loose):
                report.warnings.append(
                    f"`{IMAGES_DIR}/{loose.name}` 은 날짜 폴더 밖에 있어 건너뜁니다."
                )

    latest: dict[str, ImageEntry] = {}
    for _, folder in sorted(dated):  # oldest first; newer folders overwrite
        in_folder: dict[str, ImageEntry] = {}
        for path in sorted(folder.iterdir()):
            if not path.is_file() or _is_ignorable(path):
                continue
            reason = _unsupported_reason(path)
            if reason:
                report.warnings.append(
                    f"`{folder.name}/{path.name}` 건너뜀 — {reason}. "
                    "피그마에서 PNG(또는 JPG)로 다시 export하세요."
                )
                continue
            screen_id = _SCALE_SUFFIX_RE.sub("", path.stem).strip()
            key = _key(screen_id)
            current = in_folder.get(key)
            if current:
                keep, drop = sorted(
                    (current, ImageEntry(screen_id, path, folder.name)),
                    key=lambda e: IMAGE_EXTS.index(e.path.suffix.lower()),
                )
                report.warnings.append(
                    f"`{folder.name}` 폴더에 `{screen_id}` 파일이 여러 개 — "
                    f"`{keep.path.name}` 사용, `{drop.path.name}` 무시."
                )
                in_folder[key] = keep
            else:
                in_folder[key] = ImageEntry(screen_id, path, folder.name)
        latest.update(in_folder)
    return latest


def _image_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as img:
        return img.size


def _effective_ppi(size: tuple[int, int], region: Box, placement: str) -> float:
    """Pixels per inch of the image as shown on the slide (quality check only)."""
    per_width, per_height = size[0] / region.width, size[1] / region.height
    if placement == PLACE_TOP:
        return per_width  # width always matches the region
    # center shows the whole image (tighter axis sets the scale); fill covers the box.
    return max(per_width, per_height) if placement == PLACE_CENTER else min(per_width, per_height)


def _is_similar_ratio(ratio: float, region: Box) -> bool:
    return abs(ratio / (region.width / region.height) - 1) <= SIMILAR_RATIO_TOLERANCE


def _top_box(region: Box, ratio: float) -> tuple[Box, float]:
    """Match the region's width and top edge; return the box and bottom crop fraction.

    Top-left and top-right corners coincide with the region, so callouts placed
    on the previous screen keep their position. A slightly taller image is
    cropped at the bottom instead of spilling out of the region.
    """
    height = region.width / ratio
    crop_bottom = 0.0
    if height > region.height:
        crop_bottom = 1 - region.height / height
        height = region.height
    return Box(region.left, region.top, region.width, height), crop_bottom


def _center_box(region: Box, ratio: float) -> Box:
    if ratio > region.width / region.height:
        width, height = region.width, region.width / ratio
    else:
        width, height = region.height * ratio, region.height
    return Box(
        region.left + (region.width - width) / 2,
        region.top + (region.height - height) / 2,
        width,
        height,
    )


def _apply_fill_crop(picture, region: Box, ratio: float) -> None:
    region_ratio = region.width / region.height
    if ratio > region_ratio:
        side = (1 - region_ratio / ratio) / 2
        picture.crop_left = picture.crop_right = side
    elif ratio < region_ratio:
        side = (1 - ratio / region_ratio) / 2
        picture.crop_top = picture.crop_bottom = side


def place_image(slide, target: Target, image: ImageEntry, mode: str) -> tuple[str, float]:
    """Add the picture at the target's z-position and box, then remove the target.

    In fit mode a similar aspect ratio is top-aligned at full width; a clearly
    different one (popup, mobile) is centered. Returns the placement used and
    the effective resolution (ppi) of the placed image.
    """
    size = _image_size(image.path)
    ratio = size[0] / size[1]
    region = target.box
    crop_bottom = 0.0
    if mode == MODE_FILL:
        placement, box = PLACE_FILL, region
    elif _is_similar_ratio(ratio, region):
        placement = PLACE_TOP
        box, crop_bottom = _top_box(region, ratio)
    else:
        placement, box = PLACE_CENTER, _center_box(region, ratio)
    picture = slide.shapes.add_picture(
        str(image.path),
        Inches(box.left),
        Inches(box.top),
        Inches(box.width),
        Inches(box.height),
    )
    if mode == MODE_FILL:
        _apply_fill_crop(picture, region, ratio)
    elif crop_bottom:
        picture.crop_bottom = crop_bottom
    picture.name = f"{PIC_NAME_PREFIX}{target.screen_id}"
    picture._element.nvPicPr.cNvPr.set(
        "descr",
        f"{REGION_TAG}{region.left:.4f},{region.top:.4f},{region.width:.4f},{region.height:.4f}",
    )

    old = target.element
    old.addprevious(picture._element)
    old_rids = old.xpath(".//@r:embed")
    old.getparent().remove(old)
    # python-pptx's drop_rel() counts only r:id, not r:embed, so check usage here.
    # A re-added identical image is deduplicated onto the same rId and must survive.
    for rid in set(old_rids):
        if not slide._element.xpath(f'.//@r:embed[. = "{rid}"]'):
            slide.part.rels.pop(rid)
    return placement, _effective_ppi(size, region, placement)


def run_insert(prs, images: dict[str, ImageEntry], mode: str, report: RunReport) -> None:
    targets = find_targets(prs, report)
    used: set[str] = set()
    for target in targets:
        image = images.get(_key(target.screen_id))
        if image is None:
            if target.kind == "marker":
                report.missing_image.append(target)
            elif target.kind == "table":
                report.kept_existing.append(target)
            continue
        placement, ppi = place_image(prs.slides[target.slide_no - 1], target, image, mode)
        if ppi < MIN_PPI:
            width, height = _image_size(image.path)
            report.warnings.append(
                f"슬라이드 {target.slide_no}: `{image.path.name}` 해상도가 낮습니다 "
                f"({width}×{height}px, 표시 크기 기준 {ppi:.0f}ppi — {MIN_PPI}ppi 이상 권장). "
                "피그마에서 배율 2x로 다시 export하세요."
            )
        report.inserted.append((target, image, placement))
        used.add(_key(target.screen_id))
    report.unused_images = [img for key, img in sorted(images.items()) if key not in used]


def _natural_key(text: str) -> list:
    return [int(part) if part.isdigit() else part.upper() for part in _NUMBER_SPLIT_RE.split(text)]


def _id_table(slide):
    for shape in slide.shapes:
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                if any(_label(c.text) == TABLE_ID_LABEL for c in row.cells):
                    return shape.table
    return None


def _label(text: str) -> str:
    return re.sub(r"\s+", "", text).upper()


def _set_text(text_frame, value: str) -> None:
    """Replace the text while keeping the first run's formatting."""
    paragraphs = text_frame.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = paragraphs[0].runs
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        paragraphs[0].text = value


def _in_footer(shape, slide_height: float) -> bool:
    return _emu_to_in(shape.top) >= slide_height * FOOTER_ZONE


def _page_number_frames(slide, page_no: int, slide_height: float) -> list:
    """Text frames that show this slide's page number (footer box, "페이지" table cell)."""
    wanted = str(page_no)
    frames = []
    for shape in slide.shapes:
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = list(row.cells)
                for idx, cell in enumerate(cells[:-1]):
                    if _label(cell.text) == PAGE_LABEL and cells[idx + 1].text.strip() == wanted:
                        frames.append(cells[idx + 1].text_frame)
        elif (
            getattr(shape, "has_text_frame", False)
            and shape.text_frame.text.strip() == wanted
            and _in_footer(shape, slide_height)
        ):
            frames.append(shape.text_frame)
    return frames


def _screen_box(prs) -> Optional[Box]:
    """Most common box of the existing screenshot across 화면ID slides."""
    counts: dict[tuple, int] = {}
    for slide in prs.slides:
        picture = _largest_picture(slide) if _table_screen_id(slide) else None
        if picture is not None:
            box = _shape_box(picture)
            key = tuple(round(v, 2) for v in (box.left, box.top, box.width, box.height))
            counts[key] = counts.get(key, 0) + 1
    if not counts:
        return None
    return Box(*max(counts, key=counts.get))


def _is_skeleton(shape, page_no: int, slide_area: float, slide_height: float) -> bool:
    """Shapes every screen page shares: tables, logos, page number, the content frame."""
    if getattr(shape, "has_table", False) and shape.has_table:
        return True
    area = _emu_to_in(shape.width) * _emu_to_in(shape.height)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return area < MIN_SCREEN_AREA_SQIN
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP or not getattr(shape, "has_text_frame", False):
        return False
    text = shape.text_frame.text.strip()
    if text == str(page_no) and _in_footer(shape, slide_height):
        return True
    return shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not text and area >= slide_area * FRAME_MIN_SHARE


def _copy_rels(element, src_part, dst_part) -> None:
    for node in element.iter():
        for attr in (f"{{{_R_NS}}}embed", f"{{{_R_NS}}}link", f"{{{_R_NS}}}id"):
            rid = node.get(attr)
            if not rid or rid not in src_part.rels:
                continue
            rel = src_part.rels[rid]
            if rel.is_external:
                new_rid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, new_rid)


def _type_value(slide) -> str:
    table = _id_table(slide)
    if table is None:
        return ""
    for row in table.rows:
        cells = list(row.cells)
        for idx, cell in enumerate(cells[:-1]):
            if _label(cell.text) == TYPE_LABEL:
                return cells[idx + 1].text.strip()
    return ""


def _clone_source_index(prs, screen_indexes: list[int], position: int) -> int:
    """Screen slide nearest to ``position`` whose 유형 is "변경"; the neighbor if none.

    Pages marked 유지 use a gray frame in these decks, which a new screen must not inherit.
    """
    changed = [i for i in screen_indexes if _type_value(prs.slides[i]) == CLONE_TYPE_VALUE]
    pool = changed or screen_indexes
    return min(pool, key=lambda i: (abs(i - (position - 0.5)), i))


def _clone_screen_slide(prs, src, src_page: int, position: int):
    """Add a slide at 0-based ``position`` holding only the shared page skeleton of ``src``."""
    slide = prs.slides.add_slide(src.slide_layout)
    tree = slide.shapes._spTree
    for element in list(tree)[2:]:  # keep nvGrpSpPr / grpSpPr
        tree.remove(element)
    src_bg = src._element.cSld.bg
    if src_bg is not None:
        slide._element.cSld.insert(0, copy.deepcopy(src_bg))
    slide_height = _emu_to_in(prs.slide_height)
    slide_area = _emu_to_in(prs.slide_width) * slide_height
    for shape in src.shapes:
        if _is_skeleton(shape, src_page, slide_area, slide_height):
            element = copy.deepcopy(shape._element)
            _copy_rels(element, src.part, slide.part)
            tree.append(element)
    id_list = prs.slides._sldIdLst
    entry = id_list[-1]
    id_list.remove(entry)
    id_list.insert(position, entry)
    return slide


def _fill_new_slide(slide, screen_id: str, page_no: int, src_page: int, screen: Box, prs) -> None:
    table = _id_table(slide)
    for row in table.rows:
        cells = list(row.cells)
        for idx, cell in enumerate(cells[:-1]):
            label = _label(cell.text)
            if label == TABLE_ID_LABEL:
                _set_text(cells[idx + 1].text_frame, screen_id)
            elif label in UNKNOWN_FIELD_LABELS:
                _set_text(cells[idx + 1].text_frame, PLACEHOLDER)
    for frame in _page_number_frames(slide, src_page, _emu_to_in(prs.slide_height)):
        _set_text(frame, str(page_no))

    frame_right = _emu_to_in(prs.slide_width) - 0.3
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            frame_right = _emu_to_in(shape.left + shape.width) - 0.1
    left = screen.left + screen.width + 0.3
    note = slide.shapes.add_textbox(
        Inches(left), Inches(screen.top), Inches(max(frame_right - left, 1.0)), Inches(screen.height)
    )
    note.name = "FIGMA-NOTE"
    note.text_frame.word_wrap = True
    for idx, line in enumerate(DESCRIPTION_PLACEHOLDER):
        paragraph = note.text_frame.paragraphs[0] if idx == 0 else note.text_frame.add_paragraph()
        paragraph.text = line
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.size = Pt(10)

    marker = slide.shapes.add_textbox(
        Inches(screen.left), Inches(screen.top), Inches(screen.width), Inches(screen.height)
    )
    marker.text_frame.text = f"[화면 교체 위치: {screen_id}]"


def add_missing_slides(prs, images: dict[str, ImageEntry], report: RunReport) -> list[str]:
    """Create a screen page for each image whose ID has no slide yet.

    The new page clones the shared skeleton of the neighboring 화면ID slide,
    placed in natural ID order, and carries a marker for the regular insert
    pass. Page numbers of later slides are shifted to match.
    """
    present = {_key(t.screen_id) for t in find_targets(prs, RunReport())}
    new_ids = sorted(
        (img.screen_id for key, img in images.items() if key not in present), key=_natural_key
    )
    if not new_ids:
        return []
    screen = _screen_box(prs)
    if screen is None:
        report.warnings.append(
            "화면ID 표가 있는 슬라이드가 없어 새 슬라이드를 만들지 않았습니다 "
            f"({', '.join(new_ids)})."
        )
        return []

    slide_height = _emu_to_in(prs.slide_height)
    numbered = {
        slide.slide_id: _page_number_frames(slide, page, slide_height)
        for page, slide in enumerate(prs.slides, start=1)
    }
    added: list[str] = []
    for screen_id in new_ids:
        screens = [(i, _table_screen_id(s)) for i, s in enumerate(prs.slides) if _table_screen_id(s)]
        before = [i for i, sid in screens if _natural_key(sid) < _natural_key(screen_id)]
        neighbor = before[-1] if before else screens[0][0]
        position = neighbor + 1 if before else neighbor
        src_index = _clone_source_index(prs, [i for i, _ in screens], position)
        src = prs.slides[src_index]
        src_page = src_index + 1  # page numbers are checked against the slide position
        slide = _clone_screen_slide(prs, src, src_page, position)
        _fill_new_slide(slide, screen_id, position + 1, src_page, screen, prs)
        added.append(screen_id)

    for page, slide in enumerate(prs.slides, start=1):
        for frame in numbered.get(slide.slide_id, []):
            _set_text(frame, str(page))
    report.warnings.append(
        "새 슬라이드를 추가했습니다 — 목차·개정 이력 등 다른 페이지의 화면 목록은 자동으로 바뀌지 않으니 확인하세요."
    )
    return added


def find_template(template_dir: Path) -> Path:
    candidates = [
        p for p in sorted(template_dir.glob("*.pptx")) if not p.name.startswith("~$")
    ]
    if len(candidates) != 1:
        found = ", ".join(p.name for p in candidates) or "없음"
        raise RuntimeError(
            f"`{TEMPLATE_DIR}` 폴더에는 .pptx 파일이 정확히 1개 있어야 합니다 (현재: {found})."
        )
    return candidates[0]


def _unique_path(path: Path) -> Path:
    candidate, n = path, 2
    while candidate.exists():
        candidate = path.with_name(f"{path.stem}_{n}{path.suffix}")
        n += 1
    return candidate


def save_with_archive(prs, template: Path, output_dir: Path) -> tuple[Path, Optional[Path], list[Path]]:
    """Save the new final deck; move the previous one to _archive and prune old archives."""
    output_dir.mkdir(parents=True, exist_ok=True)
    final_path = output_dir / f"{template.stem}{OUTPUT_SUFFIX}.pptx"
    tmp_path = output_dir / f".{template.stem}{OUTPUT_SUFFIX}.tmp.pptx"
    prs.save(str(tmp_path))

    archived: Optional[Path] = None
    pruned: list[Path] = []
    if final_path.exists():
        archive_dir = output_dir / ARCHIVE_DIR
        archive_dir.mkdir(exist_ok=True)
        stamp = datetime.fromtimestamp(final_path.stat().st_mtime).strftime("%Y%m%d_%H%M")
        archived = _unique_path(archive_dir / f"{template.stem}_{stamp}.pptx")
        # Same-volume rename: fails cleanly if PowerPoint holds the file, unlike
        # shutil.move's copy+delete fallback that leaves a stray archive copy.
        try:
            final_path.rename(archived)
        except PermissionError:
            tmp_path.unlink(missing_ok=True)
            raise
        old = sorted(archive_dir.glob("*.pptx"), key=lambda p: p.stat().st_mtime, reverse=True)
        for path in old[ARCHIVE_KEEP:]:
            path.unlink()
            pruned.append(path)
    tmp_path.replace(final_path)
    return final_path, archived, pruned


def _rel(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return str(path)


def format_report(
    report: RunReport,
    *,
    added: list[str],
    root: Path,
    template: Path,
    mode: str,
    started: datetime,
    final_path: Optional[Path],
    archived: Optional[Path],
    pruned: list[Path],
) -> str:
    mode_label = "비율 유지 맞춤(fit)" if mode == MODE_FIT else "영역 채우기·잘림(fill)"
    lines = [
        f"## {started:%Y-%m-%d %H:%M} 실행",
        "",
        f"- 프로젝트: `{root.name}`",
        f"- 템플릿: `{_rel(template, root)}`",
        f"- 배치 방식: {mode_label}",
        f"- 결과 파일: `{_rel(final_path, root)}`" if final_path else "- 결과 파일: (미리보기 실행 — 저장 안 함)",
    ]
    if archived:
        lines.append(f"- 이전 최종본 보관: `{_rel(archived, root)}`")
    for path in pruned:
        lines.append(f"- 보관 개수 초과로 삭제: `{_rel(path, root)}`")
    lines += ["", f"### 삽입 성공 ({len(report.inserted)})", ""]
    if report.inserted:
        lines += ["| 슬라이드 | 화면 ID | 사용 이미지 | 대상 | 배치 |", "|---|---|---|---|---|"]
        for target, image, placement in report.inserted:
            kind = {"marker": "마커", "picture": "삽입 이미지 교체", "table": "화면ID 표·기존 화면 교체"}[
                target.kind
            ]
            lines.append(
                f"| {target.slide_no} | {target.screen_id} | `{image.folder}/{image.path.name}` | {kind} | {placement} |"
            )
    else:
        lines.append("- 없음")
    if added:
        new_pages = [t for t, _, _ in report.inserted if _key(t.screen_id) in {_key(a) for a in added}]
        lines += ["", f"### 새 슬라이드 추가 ({len(added)})", ""]
        lines += [
            f"- 슬라이드 {t.slide_no}: `{t.screen_id}` — 화면명·화면타입·유형·화면 경로·설명 {PLACEHOLDER}"
            for t in new_pages
        ]
    lines += ["", f"### 매칭 실패 — 마커는 있는데 이미지 없음 ({len(report.missing_image)})", ""]
    lines += [f"- 슬라이드 {t.slide_no}: `{t.screen_id}`" for t in report.missing_image] or ["- 없음"]
    lines += ["", f"### 이미지 없음 — 화면ID 슬라이드, 기존 화면 유지 ({len(report.kept_existing)})", ""]
    lines += [f"- 슬라이드 {t.slide_no}: `{t.screen_id}`" for t in report.kept_existing] or ["- 없음"]
    lines += ["", f"### 매칭 실패 — 이미지는 있는데 마커·화면ID 없음 ({len(report.unused_images)})", ""]
    lines += [f"- `{i.folder}/{i.path.name}`" for i in report.unused_images] or ["- 없음"]
    if report.warnings:
        lines += ["", f"### 확인 필요 ({len(report.warnings)})", ""]
        lines += [f"- {w}" for w in report.warnings]
    return "\n".join(lines) + "\n"


def append_log(log_path: Path, entry: str) -> None:
    if not log_path.exists():
        log_path.write_text(
            "# 피그마 화면 삽입 실행 이력\n\n최근 실행이 맨 아래에 추가됩니다.\n",
            encoding="utf-8",
        )
    with log_path.open("a", encoding="utf-8", newline="\n") as fh:
        fh.write("\n" + entry)


def configure_utf8_stdio() -> None:
    """Force UTF-8 console output so Korean status text survives non-UTF-8 Windows locales."""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def list_projects() -> list[str]:
    if not PROJECTS_DIR.is_dir():
        return []
    return sorted(p.name for p in PROJECTS_DIR.iterdir() if p.is_dir() and not p.name.startswith("."))


def create_project(name: str) -> Path:
    """Create the folder skeleton for a new project and return its root."""
    clean = name.strip()
    if not clean or clean.startswith(".") or clean.endswith(".") or _BAD_PROJECT_CHARS_RE.search(clean):
        raise RuntimeError(
            f"프로젝트 이름 `{name}`은 쓸 수 없습니다. "
            '\\ / : * ? " < > | 문자와 앞뒤 점(.)은 빼 주세요.'
        )
    root = PROJECTS_DIR / clean
    if root.exists():
        raise RuntimeError(f"`projects/{clean}` 프로젝트가 이미 있습니다.")
    for sub in (TEMPLATE_DIR, IMAGES_DIR, OUTPUT_DIR):
        (root / sub).mkdir(parents=True)
    return root


def resolve_project(name: Optional[str]) -> Path:
    """Return the project root for --project, or the only project when omitted."""
    projects = list_projects()
    available = ", ".join(projects) or "없음"
    if name:
        root = PROJECTS_DIR / name.strip()
        if not root.is_dir():
            raise RuntimeError(
                f"`projects/{name}` 프로젝트가 없습니다 (현재 프로젝트: {available}). "
                "새로 만들려면 --new-project를 쓰세요."
            )
        return root
    if len(projects) == 1:
        return PROJECTS_DIR / projects[0]
    if not projects:
        raise RuntimeError("프로젝트가 없습니다. --new-project <이름>으로 먼저 만드세요.")
    raise RuntimeError(f"프로젝트를 --project로 지정하세요 (현재 프로젝트: {available}).")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Insert Figma-exported screens into PPT marker shapes.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    target = parser.add_mutually_exclusive_group()
    target.add_argument(
        "--project",
        help="project folder name under projects/ (optional when only one exists)",
    )
    target.add_argument(
        "--new-project",
        metavar="NAME",
        help="create projects/NAME with empty 01_template / 02_images / 03_output and exit",
    )
    target.add_argument("--list", action="store_true", help="list projects and exit")
    target.add_argument(
        "--root",
        type=Path,
        help="use this folder directly as the project root (for testing)",
    )
    parser.add_argument(
        "--mode",
        choices=(MODE_FIT, MODE_FILL),
        default=MODE_FIT,
        help="fit: keep ratio inside the box (default); fill: cover the box and crop",
    )
    parser.add_argument(
        "--add-slides",
        action="store_true",
        help="add a screen page (cloned layout) for images whose ID has no slide yet",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="report matches only; do not save, archive, or log",
    )
    return parser


def main(argv: Optional[list[str]] = None) -> int:
    configure_utf8_stdio()
    args = build_parser().parse_args(argv)
    if args.list:
        print("\n".join(list_projects()) or "(프로젝트 없음)")
        return 0
    try:
        if args.new_project:
            root = create_project(args.new_project)
            print(
                f"`projects/{root.name}` 프로젝트를 만들었습니다.\n"
                f"- 원본 PPT 1개 → projects/{root.name}/{TEMPLATE_DIR}/\n"
                f"- 피그마 이미지 → projects/{root.name}/{IMAGES_DIR}/YYYY-MM-DD/"
            )
            return 0
        root = args.root.resolve() if args.root else resolve_project(args.project)
        template = find_template(root / TEMPLATE_DIR)
    except RuntimeError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1
    started = datetime.now()

    report = RunReport()
    images = collect_images(root / IMAGES_DIR, report)
    prs = Presentation(str(template))
    added = add_missing_slides(prs, images, report) if args.add_slides else []
    run_insert(prs, images, args.mode, report)

    final_path: Optional[Path] = None
    archived: Optional[Path] = None
    pruned: list[Path] = []
    if not args.dry_run:
        try:
            final_path, archived, pruned = save_with_archive(prs, template, root / OUTPUT_DIR)
        except PermissionError as exc:
            print(
                f"오류: 파일을 쓸 수 없습니다 ({exc.filename}). "
                "PowerPoint에서 결과 파일을 닫고 다시 실행하세요.",
                file=sys.stderr,
            )
            return 1

    entry = format_report(
        report,
        added=added,
        root=root,
        template=template,
        mode=args.mode,
        started=started,
        final_path=final_path,
        archived=archived,
        pruned=pruned,
    )
    if not args.dry_run:
        append_log(root / LOG_FILE, entry)
    print(entry)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
